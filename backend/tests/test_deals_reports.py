"""The money maths behind the auditor's sales register and the Insights tab:
IST month boundaries, the optional GST split, refunds as negative rows."""
import io
from unittest.mock import patch

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.services import deals_reports
from app.services.deals_reports import _tax_split, month_bounds
from tests.fake_supabase import FakeSupabase


def test_month_bounds_are_ist_calendar_month_in_utc():
    start, end = month_bounds("2026-09")
    # 1 Sep 00:00 IST is 31 Aug 18:30 UTC
    assert start.startswith("2026-08-31T18:30:00")
    assert end.startswith("2026-09-30T18:30:00")


def test_month_bounds_december_rolls_into_next_year():
    _, end = month_bounds("2026-12")
    assert end.startswith("2026-12-31T18:30:00")


@pytest.mark.parametrize("bad", ["2026-13", "26-09", "", "2026-9", "abc"])
def test_month_bounds_rejects_bad_input(bad):
    with pytest.raises(HTTPException) as exc:
        month_bounds(bad)
    assert exc.value.status_code == 400


def test_tax_split_price_includes_gst():
    # ₹499.00 at 18% inclusive -> taxable 422.88, tax 76.12
    split = _tax_split(49900, 18, prices_include_gst=True)
    assert split["taxable"] == 42288
    assert split["tax"] == 7612
    assert split["cgst"] + split["sgst"] == split["tax"]
    assert split["gross"] == 49900


def test_tax_split_price_excludes_gst_adds_on_top():
    split = _tax_split(49900, 18, prices_include_gst=False)
    assert split["taxable"] == 49900
    assert split["tax"] == 8982
    assert split["gross"] == 58882


def test_tax_split_odd_paisa_goes_to_cgst():
    split = _tax_split(10001, 5, prices_include_gst=False)  # tax 500 -> even; use one that is odd
    odd = _tax_split(10100, 5, prices_include_gst=False)    # tax 505
    assert odd["cgst"] == 253 and odd["sgst"] == 252
    assert split["cgst"] + split["sgst"] == split["tax"]


def test_tax_split_without_rate_is_untaxed():
    assert _tax_split(1000, None, True) == {"taxable": 1000, "tax": 0, "cgst": 0, "sgst": 0, "gross": 1000}


def _deal(number, when, items, source="walk_in", method="cash"):
    return {
        "deal_number": number, "won_at": when, "lost_at": when, "source": source, "payment_method": method,
        "razorpay_payment_id": None, "total_paise": sum(i["line_total_paise"] for i in items),
        "leads": {"name": "Ravi", "phone": "+919800000001"}, "deal_items": items,
    }


def _line(name, qty, unit, rate=None):
    return {"name": name, "qty": qty, "unit_price_paise": unit, "gst_rate": rate, "line_total_paise": qty * unit}


PROFILE = {
    "legal_name": "Kumar Mobiles", "address": "12 Anna Salai", "city": "Chennai", "state": "Tamil Nadu",
    "pincode": "600002", "gstin": "", "email": "", "phone": "9800000000", "prices_include_gst": True,
}


def _register(won, refunds, profile, fmt="xlsx"):
    with patch.object(deals_reports, "get_business_details", return_value=profile), \
         patch.object(deals_reports, "_won_in_month", return_value=won), \
         patch.object(deals_reports, "_refunded_in_month", return_value=refunds):
        return deals_reports.sales_register("tenant", "2026-09", fmt, db=object())


def test_register_header_is_the_clients_business_not_aira():
    content, media_type, filename = _register([_deal(41, "2026-09-03T05:00:00+00:00", [_line("Earbuds", 2, 129900)])], [], PROFILE)
    ws = load_workbook(io.BytesIO(content)).active
    assert ws["A1"].value == "Kumar Mobiles"
    assert "Aira" not in " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert filename == "kumar-mobiles-sales-2026-09.xlsx"
    assert media_type.endswith("spreadsheetml.sheet")


def test_register_total_nets_refunds_as_negative_rows():
    won = [_deal(41, "2026-09-03T05:00:00+00:00", [_line("Earbuds", 2, 129900)])]
    refunds = [_deal(38, "2026-09-10T05:00:00+00:00", [_line("Charger", 1, 49900)])]
    content, _, _ = _register(won, refunds, PROFILE, fmt="csv")
    text = content.decode("utf-8-sig")
    assert "D-0038 (refund)" in text
    assert "TOTAL,2099.0" in text  # 2598 - 499


def test_register_gst_columns_only_with_gstin_and_rates():
    won = [_deal(41, "2026-09-03T05:00:00+00:00", [_line("Earbuds", 1, 49900, rate=18)])]
    plain, _, _ = _register(won, [], PROFILE, fmt="csv")
    assert "CGST" not in plain.decode("utf-8-sig")
    with_gst, _, _ = _register(won, [], {**PROFILE, "gstin": "33ABCDE1234F1Z5"}, fmt="csv")
    text = with_gst.decode("utf-8-sig")
    assert "GSTIN: 33ABCDE1234F1Z5" in text
    assert "Taxable value (₹),GST %,CGST (₹),SGST (₹),IGST (₹),Amount (₹)" in text
    assert "422.88" in text


def test_register_neutralises_formula_injection_in_customer_names():
    deal = _deal(41, "2026-09-03T05:00:00+00:00", [_line("=cmd()", 1, 100)])
    deal["leads"] = {"name": "=HYPERLINK(evil)", "phone": "+91"}
    content, _, _ = _register([deal], [], PROFILE, fmt="csv")
    text = content.decode("utf-8-sig")
    assert "'=HYPERLINK(evil)" in text and "'=cmd()" in text


# ---------------------------------------------------------------- intake GST (always tax-exclusive lines)

def _intake_gst_deal():
    deal = _deal(51, "2026-09-03T05:00:00+00:00", [_line("Horoscope", 1, 4900, rate=18)], source="form", method="razorpay")
    return {**deal, "intake_session_id": "s-1", "total_paise": 5782}


@pytest.mark.parametrize("prices_include_gst", [True, False])
def test_intake_gst_line_reports_taxable_tax_and_gross_whatever_the_tenant_flag(prices_include_gst):
    profile = {**PROFILE, "gstin": "33ABCDE1234F1Z5", "prices_include_gst": prices_include_gst}
    content, _, _ = _register([_intake_gst_deal()], [], profile, fmt="csv")
    text = content.decode("utf-8-sig")
    # Rate 49, taxable 49, 18%, CGST 4.41, SGST 4.41, IGST 0, gross 57.82
    assert ",49.0,18,4.41,4.41,0,57.82," in text
    assert "TOTAL,57.82" in text


def test_tax_split_tax_exclusive_flag_overrides_the_tenant_default():
    split = _tax_split(4900, 18, prices_include_gst=True, tax_exclusive=True)
    assert (split["taxable"], split["tax"], split["cgst"], split["sgst"], split["gross"]) == (4900, 882, 441, 441, 5782)


def test_non_intake_deal_is_still_split_by_the_tenant_flag():
    profile = {**PROFILE, "gstin": "33ABCDE1234F1Z5", "prices_include_gst": True}
    won = [_deal(41, "2026-09-03T05:00:00+00:00", [_line("Earbuds", 1, 49900, rate=18)])]
    text = _register(won, [], profile, fmt="csv")[0].decode("utf-8-sig")
    assert ",422.88,18,38.06,38.06,0,499.0," in text


def test_intake_gst_refund_row_is_negative_gross():
    profile = {**PROFILE, "gstin": "33ABCDE1234F1Z5"}
    text = _register([], [_intake_gst_deal()], profile, fmt="csv")[0].decode("utf-8-sig")
    assert ",-49.0,18,-4.41,-4.41,0,-57.82," in text


def test_intake_gst_without_gstin_still_reports_what_was_charged():
    text = _register([_intake_gst_deal()], [], PROFILE, fmt="csv")[0].decode("utf-8-sig")
    assert "TOTAL,57.82" in text


# ---------------------------------------------------------------- monthly stats (Insights tab)

def _stats_db():
    db = FakeSupabase()
    for source, when, total, stage in [
        ("whatsapp", "2026-10-01T05:00:00+00:00", 1000, "won"),
        ("whatsapp", "2026-10-01T06:00:00+00:00", 500, "won"),
        ("form", "2026-10-01T07:00:00+00:00", 200, "won"),
        ("form", "2026-10-08T07:00:00+00:00", 300, "won"),
        ("form", "2026-09-30T19:00:00+00:00", 700, "won"),  # 00:30 IST Oct 1 -- October in IST
        ("whatsapp", "2026-09-10T05:00:00+00:00", 9999, "won"),  # September: out of range
    ]:
        db.add("deals", tenant_id="t1", source=source, won_at=when, total_paise=total, stage=stage, deal_items=[])
    db.add("deals", tenant_id="t1", source="whatsapp", stage="lost", lost_at="2026-10-05T05:00:00+00:00", total_paise=400)
    db.add("deals", tenant_id="t1", source="form", stage="lost", lost_at="2026-10-06T05:00:00+00:00", total_paise=400)
    db.add("deals", tenant_id="t1", source="whatsapp", stage="quoted", total_paise=800)
    db.add("deals", tenant_id="t1", source="form", stage="awaiting_payment", total_paise=600)
    db.add("deals", tenant_id="t2", source="form", won_at="2026-10-02T05:00:00+00:00", total_paise=5, stage="won", deal_items=[])
    return db


def _stats(db, sources=None):
    with patch.object(deals_reports, "_low_stock", return_value=[{"id": "x"}]):
        return deals_reports.monthly_stats("t1", "2026-10", sources=sources, db=db)


def test_stats_by_day_lists_only_days_with_sales():
    # The frontend draws every day of the month itself; the API sends sale days only.
    stats = _stats(_stats_db())
    assert [d["date"] for d in stats["by_day"]] == ["2026-10-01", "2026-10-08"]


def test_stats_by_day_splits_each_day_by_source_and_the_parts_add_up():
    day1 = _stats(_stats_db())["by_day"][0]
    assert day1["by_source"] == {"whatsapp": {"count": 2, "total_paise": 1500}, "form": {"count": 2, "total_paise": 900}}
    assert sum(v["total_paise"] for v in day1["by_source"].values()) == day1["total_paise"]
    assert sum(v["count"] for v in day1["by_source"].values()) == day1["count"]


def test_stats_source_filter_narrows_won_lost_and_open_but_not_stock():
    stats = _stats(_stats_db(), sources=["form"])
    assert stats["won_count"] == 3 and stats["won_total_paise"] == 1200
    assert stats["lost_count"] == 1
    assert stats["open_count"] == 1 and stats["open_total_paise"] == 600
    assert set(stats["by_source"]) == {"form"}
    assert stats["low_stock"] == [{"id": "x"}]


def test_stats_without_filter_counts_every_source_for_this_tenant_only():
    stats = _stats(_stats_db())
    assert stats["won_count"] == 5 and stats["won_total_paise"] == 2700
    assert stats["lost_count"] == 2 and stats["open_count"] == 2


def test_stats_route_passes_the_source_list_and_the_callers_tenant():
    from app.routes import deals as deals_routes
    with patch.object(deals_reports, "monthly_stats", return_value={}) as stats:
        deals_routes.deal_stats("2026-10", source="form,whatsapp", ctx={"tenant_id": "t1"})
        deals_routes.deal_stats("2026-10", source=None, ctx={"tenant_id": "t1"})
    assert stats.call_args_list[0].args == ("t1", "2026-10") and stats.call_args_list[0].kwargs == {"sources": ["form", "whatsapp"]}
    assert stats.call_args_list[1].kwargs == {"sources": None}


def test_stats_route_rejects_an_unknown_source():
    from app.routes import deals as deals_routes
    with pytest.raises(HTTPException) as err:
        deals_routes.deal_stats("2026-10", source="carrier_pigeon", ctx={"tenant_id": "t1"})
    assert err.value.status_code == 400


def test_active_months_route_is_scoped_to_the_callers_tenant():
    from app.routes import deals as deals_routes
    with patch.object(deals_reports, "active_months", return_value=["2026-10"]) as months:
        assert deals_routes.deal_active_months(ctx={"tenant_id": "t9"}) == {"months": ["2026-10"]}
    months.assert_called_once_with("t9")


def test_active_months_are_ist_months_with_a_won_deal_for_this_tenant_newest_first():
    assert deals_reports.active_months("t1", db=_stats_db(), today="2026-10-13") == ["2026-10", "2026-09"]
    assert deals_reports.active_months("t2", db=_stats_db(), today="2026-10-13") == ["2026-10"]
    assert deals_reports.active_months("t3", db=_stats_db(), today="2026-10-13") == []
