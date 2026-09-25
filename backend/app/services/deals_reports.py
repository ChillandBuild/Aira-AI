"""Read-only deal reports: the Insights tab numbers and the monthly sales
register a client hands to their auditor. Months are Asia/Kolkata calendar
months -- a sale at 00:30 IST on the 1st belongs to that month, not the last.

The register is the CLIENT's document: its header is their business profile
(services/business_details.py), never Aira's name."""
import csv
import io
import re
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from fastapi import HTTPException

from app.db.supabase import get_supabase
from app.services.business_details import get_business_details
from app.services.deals import LOW_STOCK_THRESHOLD, SOURCES, format_deal_number, held_quantities
from app.services.intake_csv import _csv_safe

IST = ZoneInfo("Asia/Kolkata")
MONTH_RE = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
ROW_CAP = 5000
PAYMENT_LABELS = {"razorpay": "Razorpay", "cash": "Cash", "upi": "UPI", "card": "Card", "bank_transfer": "Bank transfer", "other": "Other"}
SOURCE_LABELS = {"whatsapp": "WhatsApp", "form": "Form", "call": "Call", "walk_in": "Walk-in", "manual": "Manual", "indiamart": "IndiaMART", "justdial": "JustDial"}


def month_bounds(month: str) -> tuple[str, str]:
    match = MONTH_RE.match(month or "")
    if not match:
        raise HTTPException(status_code=400, detail="month must look like 2026-09")
    year, mon = int(match.group(1)), int(match.group(2))
    start = datetime(year, mon, 1, tzinfo=IST)
    end = datetime(year + (mon == 12), 1 if mon == 12 else mon + 1, 1, tzinfo=IST)
    return start.astimezone(timezone.utc).isoformat(), end.astimezone(timezone.utc).isoformat()


def _ist_date(iso: str) -> str:
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(IST).strftime("%Y-%m-%d")


def _won_in_month(db, tenant_id: str, start: str, end: str) -> list[dict]:
    """Deals won inside the month, whatever their stage now -- a sale refunded
    next month still happened this month."""
    return (
        db.table("deals")
        .select("*, leads(name, phone), deal_items(name, qty, unit_price_paise, gst_rate, line_total_paise)")
        .eq("tenant_id", tenant_id)
        .gte("won_at", start)
        .lt("won_at", end)
        .order("won_at")
        .limit(ROW_CAP)
        .execute()
    ).data or []


def _refunded_in_month(db, tenant_id: str, start: str, end: str) -> list[dict]:
    return (
        db.table("deals")
        .select("*, leads(name, phone), deal_items(name, qty, unit_price_paise, gst_rate, line_total_paise)")
        .eq("tenant_id", tenant_id)
        .eq("stage", "lost")
        .not_.is_("won_at", "null")
        .gte("lost_at", start)
        .lt("lost_at", end)
        .order("lost_at")
        .limit(ROW_CAP)
        .execute()
    ).data or []


# ---------------------------------------------------------------- stats

def monthly_stats(tenant_id: str, month: str, db=None) -> dict:
    db = db or get_supabase()
    start, end = month_bounds(month)
    won = _won_in_month(db, tenant_id, start, end)
    lost = (
        db.table("deals").select("id").eq("tenant_id", tenant_id).eq("stage", "lost")
        .gte("lost_at", start).lt("lost_at", end).limit(ROW_CAP).execute()
    ).data or []
    open_rows = (
        db.table("deals").select("total_paise").eq("tenant_id", tenant_id)
        .in_("stage", ["quoted", "awaiting_payment"]).limit(ROW_CAP).execute()
    ).data or []

    by_source: dict[str, dict] = {}
    by_day: dict[str, dict] = {}
    items: dict[str, dict] = {}
    for deal in won:
        src = by_source.setdefault(deal["source"], {"count": 0, "total_paise": 0})
        src["count"] += 1
        src["total_paise"] += deal["total_paise"]
        day = by_day.setdefault(_ist_date(deal["won_at"]), {"count": 0, "total_paise": 0})
        day["count"] += 1
        day["total_paise"] += deal["total_paise"]
        for line in deal.get("deal_items") or []:
            agg = items.setdefault(line["name"], {"name": line["name"], "qty": 0, "total_paise": 0})
            agg["qty"] += line["qty"]
            agg["total_paise"] += line["line_total_paise"]

    return {
        "month": month,
        "won_count": len(won),
        "won_total_paise": sum(d["total_paise"] for d in won),
        "lost_count": len(lost),
        "open_count": len(open_rows),
        "open_total_paise": sum(r["total_paise"] for r in open_rows),
        "by_source": {s: by_source[s] for s in SOURCES if s in by_source},
        "by_day": [{"date": d, **v} for d, v in sorted(by_day.items())],
        "top_items": sorted(items.values(), key=lambda i: i["total_paise"], reverse=True)[:10],
        "low_stock": _low_stock(db, tenant_id),
    }


def _low_stock(db, tenant_id: str) -> list[dict]:
    tracked = (
        db.table("catalog_items").select("id, name, stock_quantity").eq("tenant_id", tenant_id)
        .not_.is_("stock_quantity", "null").limit(2000).execute()
    ).data or []
    held = held_quantities(tenant_id, [i["id"] for i in tracked], db)
    low = [
        {**i, "held_quantity": held.get(i["id"], 0)}
        for i in tracked
        if i["stock_quantity"] - held.get(i["id"], 0) <= LOW_STOCK_THRESHOLD
    ]
    return sorted(low, key=lambda i: i["stock_quantity"] - i["held_quantity"])[:20]


# ---------------------------------------------------------------- sales register

def _tax_split(line_total: int, rate: float | None, prices_include_gst: bool) -> dict:
    """Intra-state assumed (customer state isn't collected): CGST = SGST, the
    odd paisa goes to CGST. IGST stays 0 but the column exists for the auditor."""
    if not rate:
        return {"taxable": line_total, "tax": 0, "cgst": 0, "sgst": 0, "gross": line_total}
    if prices_include_gst:
        taxable = round(line_total * 100 / (100 + rate))
        tax = line_total - taxable
    else:
        taxable = line_total
        tax = round(line_total * rate / 100)
    sgst = tax // 2
    return {"taxable": taxable, "tax": tax, "cgst": tax - sgst, "sgst": sgst, "gross": taxable + tax}


def _register_rows(deals: list[dict], sign: int, date_key: str, include_gst: bool, profile: dict) -> list[list]:
    rows = []
    for deal in deals:
        lead = deal.get("leads") or {}
        for line in deal.get("deal_items") or []:
            split = _tax_split(line["line_total_paise"], line.get("gst_rate"), profile["prices_include_gst"])
            row = [
                _ist_date(deal[date_key]),
                format_deal_number(deal["deal_number"]) + (" (refund)" if sign < 0 else ""),
                _csv_safe(lead.get("name") or ""),
                _csv_safe(lead.get("phone") or ""),
                _csv_safe(line["name"]),
                sign * line["qty"],
                line["unit_price_paise"] / 100,
            ]
            if include_gst:
                row += [
                    sign * split["taxable"] / 100,
                    line.get("gst_rate") or 0,
                    sign * split["cgst"] / 100,
                    sign * split["sgst"] / 100,
                    0,
                    sign * split["gross"] / 100,
                ]
            else:
                row += [sign * line["line_total_paise"] / 100]
            row += [
                PAYMENT_LABELS.get(deal.get("payment_method") or "", ""),
                _csv_safe(deal.get("razorpay_payment_id") or ""),
                SOURCE_LABELS.get(deal["source"], deal["source"]),
            ]
            rows.append(row)
    return rows


def sales_register(tenant_id: str, month: str, fmt: str, db=None) -> tuple[bytes, str, str]:
    """Returns (file bytes, media type, filename). One row per sold line;
    refunds (won, then marked lost) appear as negative rows in the month they
    were refunded, so the month total matches money actually kept."""
    db = db or get_supabase()
    start, end = month_bounds(month)
    profile = get_business_details(tenant_id, db)
    won = _won_in_month(db, tenant_id, start, end)
    refunds = _refunded_in_month(db, tenant_id, start, end)
    has_rates = any(line.get("gst_rate") for d in won + refunds for line in d.get("deal_items") or [])
    include_gst = bool(profile.get("gstin")) and has_rates

    headers = ["Date", "Deal #", "Customer", "Phone", "Item", "Qty", "Rate (₹)"]
    headers += (["Taxable value (₹)", "GST %", "CGST (₹)", "SGST (₹)", "IGST (₹)", "Amount (₹)"] if include_gst else ["Amount (₹)"])
    headers += ["Paid by", "Payment ref", "Source"]
    body = _register_rows(won, 1, "won_at", include_gst, profile) + _register_rows(refunds, -1, "lost_at", include_gst, profile)
    amount_col = headers.index("Amount (₹)")
    total = round(sum(r[amount_col] for r in body), 2)

    month_label = datetime.strptime(month, "%Y-%m").strftime("%B %Y")
    title_lines = [
        profile.get("legal_name") or "",
        ", ".join(p for p in (profile.get("address"), profile.get("city"), profile.get("state"), profile.get("pincode")) if p),
        f"GSTIN: {profile['gstin']}" if profile.get("gstin") else "",
        " · ".join(p for p in (profile.get("phone"), profile.get("email")) if p),
        f"Sales register — {month_label}",
    ]
    title_lines = [_csv_safe(t) for t in title_lines if t]
    total_row = [""] * len(headers)
    total_row[amount_col - 1] = "TOTAL"
    total_row[amount_col] = total
    filename_base = re.sub(r"[^A-Za-z0-9]+", "-", profile.get("legal_name") or "sales").strip("-").lower() or "sales"

    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        for t in title_lines:
            writer.writerow([t])
        writer.writerow([])
        writer.writerow(headers)
        writer.writerows(body)
        writer.writerow(total_row)
        return buf.getvalue().encode("utf-8-sig"), "text/csv", f"{filename_base}-sales-{month}.csv"

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = f"Sales {month}"
    for i, t in enumerate(title_lines):
        ws.append([t])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=14 if i == 0 else 11)
    ws.append([])
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for r in body:
        ws.append(r)
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHIJKLMNOP", [12, 16, 22, 16, 30, 7, 11, 16, 8, 11, 11, 10, 14, 14, 20, 12]):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return (
        out.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"{filename_base}-sales-{month}.xlsx",
    )


# ---------------------------------------------------------------- ask-box summary

def sales_summary(tenant_id: str, period: str = "today", db=None) -> dict:
    """What the dashboard's Ask box reads for sales questions ("how much did we
    sell today", "who hasn't paid", "what's low on stock"). Periods are IST."""
    db = db or get_supabase()
    now = datetime.now(IST)
    day_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    starts = {
        "today": day_start,
        "this_week": day_start - timedelta(days=day_start.weekday()),
        "this_month": day_start.replace(day=1),
    }
    period = period if period in starts else "today"
    start = starts[period].astimezone(timezone.utc).isoformat()
    won = (
        db.table("deals").select("total_paise, source, deal_items(name, qty)").eq("tenant_id", tenant_id)
        .gte("won_at", start).limit(ROW_CAP).execute()
    ).data or []
    unpaid = (
        db.table("deals").select("deal_number, total_paise, leads(name, phone)").eq("tenant_id", tenant_id)
        .eq("stage", "awaiting_payment").order("created_at", desc=True).limit(50).execute()
    ).data or []
    items: dict[str, int] = {}
    for deal in won:
        for line in deal.get("deal_items") or []:
            items[line["name"]] = items.get(line["name"], 0) + line["qty"]
    return {
        "period": period,
        "sales_count": len(won),
        "sales_total_rupees": sum(d["total_paise"] for d in won) / 100,
        "top_items_by_quantity": sorted(({"name": k, "qty": v} for k, v in items.items()), key=lambda i: -i["qty"])[:5],
        "unpaid_payment_links": [
            {"deal": format_deal_number(d["deal_number"]), "customer": (d.get("leads") or {}).get("name") or (d.get("leads") or {}).get("phone"),
             "amount_rupees": d["total_paise"] / 100}
            for d in unpaid
        ],
        "low_stock_items": [
            {"name": i["name"], "in_stock": i["stock_quantity"], "held_for_unpaid": i["held_quantity"]}
            for i in _low_stock(db, tenant_id)
        ],
    }
