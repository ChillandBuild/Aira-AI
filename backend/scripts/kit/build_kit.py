"""Build the Aira Business Kit Word templates, one per industry, into
frontend/public/kit/. Spec: docs/superpowers/specs/2026-09-24-aira-business-kit-design.md

Each file: a "how to use" box, the blank Kit with [WRITE ...] hints, then the example
business below a delete-me marker. Two things keep the invented example and the
instructions out of a client's knowledge even if they upload the file untouched:
- the "how to use" box is a table, and extract_text_from_file reads only paragraphs;
- knowledge_kit.strip_example cuts everything from the marker line on, and
  scrub_placeholders drops every unfilled [WRITE ...] hint.

Headings carry no numbers: knowledge_sections only treats a line as a heading when it
starts with a capital letter.

    cd backend && .venv/bin/python -m scripts.kit.build_kit
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from scripts.kit.examples_a import EXAMPLES_A
from scripts.kit.examples_b import EXAMPLES_B
from scripts.kit.examples_c import EXAMPLES_C

OUT_DIR = Path(__file__).resolve().parents[3] / "frontend" / "public" / "kit"
EXAMPLES = EXAMPLES_A + EXAMPLES_B + EXAMPLES_C
EXAMPLE_MARKER = "=== EXAMPLE BELOW - DELETE BEFORE UPLOADING ==="

VIOLET = RGBColor(0x5B, 0x21, 0xB6)
GREY = RGBColor(0x78, 0x71, 0x6C)

HEADINGS = [
    ("ABOUT YOUR BUSINESS",
     ["[WRITE your business name, what you do, since which year, and where you serve]"]),
    ("WHO YOUR CUSTOMERS ARE",
     ["[WRITE who usually messages you, and what they ask about most]"]),
    ("HOW A CUSTOMER BUYS FROM YOU",
     ["[WRITE the steps from first message to paying, one step per line]",
      "[WRITE Aira's goal: the one next step Aira should push every interested customer toward]"]),
    ("HOW AIRA SHOULD SOUND",
     ["[WRITE the tone, how to address customers, and how long replies should be. Optional]"]),
    ("WHAT AIRA MUST NEVER SAY OR PROMISE",
     ["[WRITE one thing per line, for example: never promise a discount]"]),
    ("WHEN TO HAND OVER TO A PERSON",
     ["[WRITE which situations need a person, plus your phone number and working hours]"]),
    ("PRODUCTS, SERVICES, PRICES",
     ["[WRITE each product or service with its exact price, what is included and what is not]"]),
    ("CUSTOMER QUESTIONS AND POLICIES",
     ["[WRITE a question customers really ask, starting with Q:]",
      "[WRITE your answer, starting with A:]",
      "[WRITE your payment, refund, cancellation and delivery rules]"]),
]

HOW_TO_USE = [
    "How to use this file",
    "1. Replace every [WRITE ...] line with your own words. Plain, short lines are best.",
    "2. Copy prices, phone numbers and links exactly.",
    "3. Not sure about something? Leave the [WRITE ...] line as it is. Aira will skip it and remind you later.",
    "4. The example at the end shows a finished Kit. Delete it before uploading. If you forget, Aira ignores it.",
    "5. Save, then upload this file on the Knowledge page in Aira.",
]


def _shade(cell, hex_fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    props.append(shading)


def _line(doc, text: str, *, bold=False, italic=False, color=None, size=11, space_after=2):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold, run.italic, run.font.size = bold, italic, Pt(size)
    if color is not None:
        run.font.color.rgb = color
    para.paragraph_format.space_after = Pt(space_after)
    return para


def _heading(doc, text: str) -> None:
    para = _line(doc, text, bold=True, color=VIOLET, size=12, space_after=4)
    para.paragraph_format.space_before = Pt(14)


def _how_to_use_box(doc) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _shade(cell, "F5F3FF")
    cell.paragraphs[0].add_run(HOW_TO_USE[0]).bold = True
    for text in HOW_TO_USE[1:]:
        cell.add_paragraph(text).paragraph_format.space_after = Pt(2)


def build(example: dict) -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    _line(doc, "Aira Business Kit", bold=True, size=20, space_after=0)
    _line(doc, f"Template for: {example['industry']}", color=GREY, space_after=10)
    _how_to_use_box(doc)

    for heading, hints in HEADINGS:
        _heading(doc, heading)
        for hint in hints:
            _line(doc, hint, italic=True, color=GREY)

    doc.add_page_break()
    _line(doc, EXAMPLE_MARKER, bold=True, color=VIOLET, size=12, space_after=4)
    _line(doc, f"Example: {example['business']}. A made-up business. Every name, price and number is invented.",
          italic=True, color=GREY, space_after=6)
    for (heading, _hints), body in zip(HEADINGS, example["sections"], strict=True):
        _heading(doc, heading)
        for text in body.split("\n"):
            _line(doc, text)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"aira-kit-{example['slug']}.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for ex in EXAMPLES:
        print(build(ex).relative_to(OUT_DIR.parents[1]))
